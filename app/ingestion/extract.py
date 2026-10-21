"""Text extraction from PDF (PyMuPDF), Word (python-docx) and plain-text files."""

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md"})


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedPage:
    page: int | None  # 1-based page number when the format has pages
    text: str


def extract(path: Path) -> list[ExtractedPage]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return [ExtractedPage(page=None, text=path.read_text(encoding="utf-8", errors="replace"))]
    raise UnsupportedDocumentError(f"unsupported file type: {suffix or '(none)'}")


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    import fitz  # PyMuPDF

    with fitz.open(path) as document:
        return [
            ExtractedPage(page=number, text=page.get_text("text"))
            for number, page in enumerate(document, start=1)
        ]


def _extract_docx(path: Path) -> list[ExtractedPage]:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [ExtractedPage(page=None, text="\n\n".join(parts))]
