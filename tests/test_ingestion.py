from pathlib import Path

import pytest
from pydantic import ValidationError

from app.ingestion.chunk import chunk_pages, chunk_text
from app.ingestion.clean import clean_text
from app.ingestion.extract import ExtractedPage, UnsupportedDocumentError, extract
from app.ingestion.metadata import DocumentMetadata, load_metadata
from app.ingestion.pipeline import IngestionError, prepare_document
from tests.conftest import HashingEmbedder

DEMO_DIR = Path(__file__).resolve().parents[1] / "data" / "demo"

GUIDE_EXAMPLE = {
    "document_id": "doc-001",
    "title": "HR Compensation Policy",
    "department": "HR",
    "allowed_groups": ["hr-team", "executives"],
    "classification": "confidential",
    "version": 3,
}


def test_clean_text_normalises_whitespace_and_hyphenation():
    raw = "Parental  leave\x00 is avail-\nable.\n\n\n\nNext \t para  "
    assert clean_text(raw) == "Parental leave is available.\n\nNext para"


def test_chunk_text_respects_size_and_overlap():
    words = [f"w{i}" for i in range(500)]
    chunks = [chunk.split() for chunk in chunk_text(" ".join(words), size=100, overlap=20)]

    assert all(len(chunk) <= 100 for chunk in chunks)
    assert chunks[0][-20:] == chunks[1][:20]
    assert set().union(*chunks) == set(words)


def test_chunk_text_packs_short_paragraphs_together():
    assert chunk_text("alpha beta\n\ngamma delta\n\nepsilon", size=10, overlap=2) == [
        "alpha beta gamma delta epsilon"
    ]


def test_chunk_text_handles_empty_text_and_rejects_bad_overlap():
    assert chunk_text(" \n\n ", size=10, overlap=2) == []
    with pytest.raises(ValueError):
        chunk_text("a b", size=10, overlap=10)


def test_chunk_pages_keeps_page_numbers_for_citations():
    chunks = chunk_pages([ExtractedPage(1, "one two"), ExtractedPage(2, "three four")], size=50, overlap=5)
    assert [(chunk.index, chunk.page) for chunk in chunks] == [(0, 1), (1, 2)]


def test_extract_text_file(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("# Title\n\nBody", encoding="utf-8")
    assert extract(path) == [ExtractedPage(None, "# Title\n\nBody")]


def test_extract_docx_includes_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Hello from Word")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Q3"
    table.cell(0, 1).text = "18.4M"
    path = tmp_path / "budget.docx"
    document.save(str(path))

    text = extract(path)[0].text
    assert "Hello from Word" in text
    assert "Q3 | 18.4M" in text


def test_extract_pdf_keeps_pages(tmp_path):
    import fitz

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Budget page one")
    pdf.new_page().insert_text((72, 72), "Second page")
    path = tmp_path / "budget.pdf"
    pdf.save(str(path))
    pdf.close()

    pages = extract(path)
    assert [page.page for page in pages] == [1, 2]
    assert "Budget page one" in pages[0].text


def test_extract_rejects_unsupported_types(tmp_path):
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(b"not really a spreadsheet")
    with pytest.raises(UnsupportedDocumentError):
        extract(path)


def test_metadata_accepts_the_guide_example():
    metadata = DocumentMetadata.model_validate(GUIDE_EXAMPLE)
    assert metadata.acl_principals() == ["group:executives", "group:hr-team"]


def test_metadata_rejects_unknown_fields():
    # A typo such as "allowed_group" must fail loudly rather than silently drop the ACL.
    with pytest.raises(ValidationError):
        DocumentMetadata.model_validate({**GUIDE_EXAMPLE, "allowed_group": ["marketing-team"]})


def test_metadata_requires_someone_to_be_allowed():
    with pytest.raises(ValidationError):
        DocumentMetadata.model_validate({**GUIDE_EXAMPLE, "allowed_groups": []})


def test_confidential_documents_cannot_be_public():
    with pytest.raises(ValidationError):
        DocumentMetadata.model_validate({**GUIDE_EXAMPLE, "public": True})


def test_demo_documents_have_valid_metadata():
    documents = sorted(DEMO_DIR.rglob("*.md"))
    assert len(documents) == 4
    ids = {load_metadata(path).document_id for path in documents}
    assert ids == {"doc-hr-001", "doc-hr-002", "doc-fin-001", "doc-mkt-001"}


def test_prepare_document_embeds_every_chunk():
    prepared = prepare_document(
        DEMO_DIR / "hr" / "parental-leave-policy.md", HashingEmbedder(), chunk_size=50, chunk_overlap=10
    )
    assert len(prepared.chunks) > 1
    assert len(prepared.embeddings) == len(prepared.chunks)
    assert all(len(vector) == 384 for vector in prepared.embeddings)


def test_prepare_document_fails_on_empty_or_unsupported_files(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n", encoding="utf-8")
    with pytest.raises(IngestionError):
        prepare_document(empty, HashingEmbedder(), chunk_size=50, chunk_overlap=10)

    other = tmp_path / "image.png"
    other.write_bytes(b"\x89PNG")
    with pytest.raises(IngestionError):
        prepare_document(other, HashingEmbedder(), chunk_size=50, chunk_overlap=10)
